#!/usr/bin/env python3
"""
Script to analyze the Excel spreadsheet and extract student data
"""
import openpyxl
import json

def analyze_spreadsheet(filename):
    """Analyze the Excel spreadsheet structure and contents"""
    try:
        workbook = openpyxl.load_workbook(filename)
        print(f"Loaded workbook with sheets: {workbook.sheetnames}")
        
        analysis = {
            'sheets': {},
            'students_data': {},
            'data_collection_structure': {}
        }
        
        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            print(f"\n--- Analyzing sheet: {sheet_name} ---")
            
            # Get sheet dimensions
            max_row = sheet.max_row
            max_col = sheet.max_column
            print(f"Sheet dimensions: {max_row} rows x {max_col} columns")
            
            # Read sheet data
            sheet_data = []
            for row in range(1, min(21, max_row + 1)):  # Read first 20 rows or all if less
                row_data = []
                for col in range(1, min(11, max_col + 1)):  # Read first 10 columns or all if less  
                    cell_value = sheet.cell(row=row, column=col).value
                    if cell_value is not None:
                        row_data.append(str(cell_value))
                    else:
                        row_data.append("")
                sheet_data.append(row_data)
            
            analysis['sheets'][sheet_name] = {
                'dimensions': f"{max_row}x{max_col}",
                'sample_data': sheet_data
            }
            
            # Special handling for specific sheets
            if sheet_name.lower() in ['database', 'db']:
                print("Found database sheet - analyzing student/objective structure...")
                students_data = extract_students_from_database_sheet(sheet)
                analysis['students_data'] = students_data
                
            elif sheet_name.lower() in ['data collection', 'datacollection', 'data_collection']:
                print("Found data collection sheet - analyzing structure...")
                data_structure = analyze_data_collection_sheet(sheet)
                analysis['data_collection_structure'] = data_structure
        
        return analysis
        
    except Exception as e:
        print(f"Error analyzing spreadsheet: {e}")
        return None

def extract_students_from_database_sheet(sheet):
    """Extract students and their objectives from the database sheet"""
    students_data = {}
    
    # Assume first row contains student names as column headers
    header_row = 1
    max_col = sheet.max_column
    max_row = sheet.max_row
    
    # Get student names from header row
    student_names = []
    for col in range(1, max_col + 1):
        cell_value = sheet.cell(row=header_row, column=col).value
        if cell_value and str(cell_value).strip():
            student_names.append((col, str(cell_value).strip()))
    
    print(f"Found potential students: {[name for _, name in student_names]}")
    
    # Extract objectives for each student
    for col, student_name in student_names:
        objectives = []
        for row in range(header_row + 1, max_row + 1):
            cell_value = sheet.cell(row=row, column=col).value
            if cell_value and str(cell_value).strip():
                objectives.append(str(cell_value).strip())
        
        if objectives:  # Only add students who have objectives
            students_data[student_name] = objectives
            print(f"{student_name}: {len(objectives)} objectives")
    
    return students_data

def analyze_data_collection_sheet(sheet):
    """Analyze the data collection sheet structure"""
    structure = {
        'headers': [],
        'sample_rows': [],
        'key_columns': {}
    }
    
    # Get headers from first row
    max_col = sheet.max_column
    for col in range(1, max_col + 1):
        header = sheet.cell(row=1, column=col).value
        if header:
            structure['headers'].append(str(header).strip())
        else:
            structure['headers'].append(f"Column_{col}")
    
    # Get sample data rows (up to 10)
    max_row = min(11, sheet.max_row)
    for row in range(2, max_row + 1):
        row_data = []
        for col in range(1, max_col + 1):
            cell_value = sheet.cell(row=row, column=col).value
            if cell_value is not None:
                row_data.append(str(cell_value))
            else:
                row_data.append("")
        structure['sample_rows'].append(row_data)
    
    # Identify key columns
    for i, header in enumerate(structure['headers']):
        header_lower = header.lower()
        if 'date' in header_lower:
            structure['key_columns']['date'] = i
        elif 'activity' in header_lower:
            structure['key_columns']['activity'] = i
        elif 'data' in header_lower:
            structure['key_columns']['data'] = i
        elif 'prompt' in header_lower:
            structure['key_columns']['prompt'] = i
    
    return structure

if __name__ == "__main__":
    filename = "attached_assets/Sam_Student_Dat_25-26_1757599841065.xlsx"
    analysis = analyze_spreadsheet(filename)
    
    if analysis:
        print(f"\n=== ANALYSIS COMPLETE ===")
        print(f"Found {len(analysis['students_data'])} students with data")
        
        # Pretty print the analysis
        with open('spreadsheet_analysis.json', 'w') as f:
            json.dump(analysis, f, indent=2)
        print("Analysis saved to spreadsheet_analysis.json")
        
        # Print summary
        if analysis['students_data']:
            print("\nStudents found:")
            for student, objectives in analysis['students_data'].items():
                print(f"  {student}: {len(objectives)} objectives")
        
        if analysis['data_collection_structure']['headers']:
            print(f"\nData collection sheet headers: {analysis['data_collection_structure']['headers']}")
            print(f"Key columns identified: {analysis['data_collection_structure']['key_columns']}")